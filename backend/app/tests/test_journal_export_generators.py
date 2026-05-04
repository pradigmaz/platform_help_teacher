import os
from datetime import UTC, date, datetime
from io import BytesIO
from uuid import uuid4

os.environ.setdefault("POSTGRES_SERVER", "localhost")
os.environ.setdefault("POSTGRES_USER", "test")
os.environ.setdefault("POSTGRES_PASSWORD", "test")
os.environ.setdefault("POSTGRES_DB", "test")
os.environ.setdefault("SECRET_KEY", "test-secret-key-for-testing-only")
os.environ.setdefault("ENVIRONMENT", "test")

from openpyxl import load_workbook

from app.schemas.export import (
    AttendanceExportRow,
    GradeExportRow,
    JournalExportData,
    JournalExportMeta,
    LessonExportColumn,
)
from app.services.export.csv_generator import generate_grades_csv
from app.services.export.excel_generator import generate_excel


def _export_data() -> JournalExportData:
    first_lesson = uuid4()
    second_lesson = uuid4()
    student_id = uuid4()
    return JournalExportData(
        meta=JournalExportMeta(
            group_code="ИС-1",
            group_name="ИС-1",
            period_start=date(2026, 2, 1),
            period_end=date(2026, 4, 30),
            generated_at=datetime(2026, 4, 27, tzinfo=UTC),
            total_students=1,
            total_lessons=2,
        ),
        lessons=[
            LessonExportColumn(
                lesson_id=first_lesson,
                date=date(2026, 2, 10),
                lesson_number=1,
                lesson_type="lab",
                topic=None,
                work_number=1,
                subgroup=None,
            ),
            LessonExportColumn(
                lesson_id=second_lesson,
                date=date(2026, 4, 14),
                lesson_number=2,
                lesson_type="lab",
                topic=None,
                work_number=2,
                subgroup=2,
            ),
        ],
        attendance_rows=[
            AttendanceExportRow(
                student_id=student_id,
                student_name="Иванов Иван",
                subgroup=1,
                attendance_by_date={"2026-02-10_1": "PRESENT"},
                stats={"present_count": 1, "absent_count": 0, "late_count": 0, "excused_count": 0, "total": 1},
                attendance_rate=100,
            )
        ],
        grade_rows=[
            GradeExportRow(
                student_id=student_id,
                student_name="Иванов Иван",
                subgroup=1,
                grades_by_work={"2026-02-10_1_1": 5},
                grades_count=1,
            )
        ],
    )


def test_grades_csv_includes_all_lesson_dates_and_no_average_column():
    csv_content = generate_grades_csv(_export_data())

    assert "10.02 №1 р1" in csv_content
    assert "14.04 №2 р2" in csv_content
    assert "Средняя" not in csv_content
    assert "Иванов Иван;1;5;;1" in csv_content


def test_grades_excel_includes_all_lesson_dates_and_no_average_values():
    workbook = load_workbook(BytesIO(generate_excel(_export_data())))
    grades_sheet = workbook["Оценки"]
    summary_sheet = workbook["Сводка"]

    headers = [grades_sheet.cell(row=1, column=column).value for column in range(1, 7)]
    assert headers == ["№", "ФИО", "Подгр.", "10.02\n№1/р1", "14.04\n№2/р2", "Кол-во"]
    assert grades_sheet.cell(row=2, column=4).value == 5
    assert grades_sheet.cell(row=2, column=5).value is None
    assert "Средняя оценка:" not in [summary_sheet.cell(row=row, column=1).value for row in range(1, 12)]


def test_excel_marks_other_subgroup_lesson_cells_gray():
    workbook = load_workbook(BytesIO(generate_excel(_export_data())))
    attendance_sheet = workbook["Посещаемость"]
    grades_sheet = workbook["Оценки"]

    assert attendance_sheet.cell(row=2, column=5).fill.fgColor.rgb.endswith("D9D9D9")
    assert grades_sheet.cell(row=2, column=5).fill.fgColor.rgb.endswith("D9D9D9")
