"""Генератор Excel файлов для экспорта журнала."""

import logging
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.schemas.export import (
    JournalExportData,
)

logger = logging.getLogger(__name__)

# Цвета для статусов посещаемости
STATUS_COLORS = {
    "PRESENT": "90EE90",  # светло-зелёный
    "LATE": "FFD700",  # жёлтый
    "EXCUSED": "87CEEB",  # голубой
    "ABSENT": "FF6B6B",  # красный
}

# Символы для статусов
STATUS_SYMBOLS = {
    "PRESENT": "✓",
    "LATE": "О",
    "EXCUSED": "У",
    "ABSENT": "Н",
}

# Стили
HEADER_FONT = Font(bold=True, size=10)
HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
CELL_ALIGNMENT = Alignment(horizontal="center", vertical="center")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def _auto_column_width(ws, min_width: int = 5, max_width: int = 30) -> None:
    """Автоматическая ширина колонок."""
    for column_cells in ws.columns:
        max_length = 0
        column = column_cells[0].column_letter
        for cell in column_cells:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except (TypeError, AttributeError):
                pass
        adjusted_width = min(max(max_length + 2, min_width), max_width)
        ws.column_dimensions[column].width = adjusted_width


def _apply_header_style(cell) -> None:
    """Применить стиль заголовка к ячейке."""
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = HEADER_ALIGNMENT
    cell.border = THIN_BORDER


def _apply_cell_style(cell, status: str | None = None) -> None:
    """Применить стиль к ячейке данных."""
    cell.alignment = CELL_ALIGNMENT
    cell.border = THIN_BORDER
    if status and status in STATUS_COLORS:
        cell.fill = PatternFill(
            start_color=STATUS_COLORS[status],
            end_color=STATUS_COLORS[status],
            fill_type="solid",
        )


def generate_attendance_sheet(wb: Workbook, data: JournalExportData) -> None:
    """
    Создаёт лист "Посещаемость".

    Args:
        wb: Рабочая книга Excel
        data: Данные журнала для экспорта
    """
    ws = wb.create_sheet("Посещаемость")

    # Заголовки: №, ФИО, Подгруппа, [занятия...], Всего, Присут., Отсут., Опозд., Уваж., %
    headers = ["№", "ФИО", "Подгр."]

    # Добавляем колонки занятий
    lesson_keys = []
    for lesson in data.lessons:
        col_header = f"{lesson.date.strftime('%d.%m')}\n№{lesson.lesson_number}"
        headers.append(col_header)
        lesson_keys.append(f"{lesson.date}_{lesson.lesson_number}")

    # Статистические колонки
    headers.extend(["Всего", "Присут.", "Отсут.", "Опозд.", "Уваж.", "%"])

    # Записываем заголовки
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        _apply_header_style(cell)

    # Записываем данные студентов
    for row_idx, row_data in enumerate(data.attendance_rows, 2):
        # №
        cell = ws.cell(row=row_idx, column=1, value=row_idx - 1)
        _apply_cell_style(cell)

        # ФИО
        cell = ws.cell(row=row_idx, column=2, value=row_data.student_name)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = THIN_BORDER

        # Подгруппа
        subgroup_val = row_data.subgroup if row_data.subgroup else "-"
        cell = ws.cell(row=row_idx, column=3, value=subgroup_val)
        _apply_cell_style(cell)

        # Посещаемость по занятиям
        col_offset = 4
        for i, key in enumerate(lesson_keys):
            status = row_data.attendance_by_date.get(key, "")
            symbol = STATUS_SYMBOLS.get(status, "")
            cell = ws.cell(row=row_idx, column=col_offset + i, value=symbol)
            _apply_cell_style(cell, status)

        # Статистика
        stats_col = col_offset + len(lesson_keys)
        stats = row_data.stats

        cell = ws.cell(row=row_idx, column=stats_col, value=stats.get("total", 0))
        _apply_cell_style(cell)

        cell = ws.cell(row=row_idx, column=stats_col + 1, value=stats.get("present_count", 0))
        _apply_cell_style(cell)

        cell = ws.cell(row=row_idx, column=stats_col + 2, value=stats.get("absent_count", 0))
        _apply_cell_style(cell)

        cell = ws.cell(row=row_idx, column=stats_col + 3, value=stats.get("late_count", 0))
        _apply_cell_style(cell)

        cell = ws.cell(row=row_idx, column=stats_col + 4, value=stats.get("excused_count", 0))
        _apply_cell_style(cell)

        cell = ws.cell(row=row_idx, column=stats_col + 5, value=f"{row_data.attendance_rate:.1f}%")
        _apply_cell_style(cell)

    _auto_column_width(ws)
    # Фиксируем ширину колонки ФИО
    ws.column_dimensions["B"].width = 25

    logger.debug("Лист 'Посещаемость' создан: %d строк", len(data.attendance_rows))


def generate_grades_sheet(wb: Workbook, data: JournalExportData) -> None:
    """
    Создаёт лист "Оценки".

    Args:
        wb: Рабочая книга Excel
        data: Данные журнала для экспорта
    """
    ws = wb.create_sheet("Оценки")

    # Заголовки: №, ФИО, Подгруппа, [работы...], Кол-во, Средняя
    headers = ["№", "ФИО", "Подгр."]

    # Собираем уникальные ключи работ из всех строк
    work_keys: set[str] = set()
    for row_data in data.grade_rows:
        work_keys.update(row_data.grades_by_work.keys())

    # Сортируем ключи по дате и номеру
    sorted_work_keys = sorted(work_keys)

    # Добавляем колонки работ
    for key in sorted_work_keys:
        parts = key.split("_")
        if len(parts) >= 2:
            date_str = parts[0]
            try:
                from datetime import datetime

                dt = datetime.strptime(date_str, "%Y-%m-%d")
                col_header = f"{dt.strftime('%d.%m')}"
            except ValueError:
                col_header = key
        else:
            col_header = key
        headers.append(col_header)

    headers.extend(["Кол-во", "Средняя"])

    # Записываем заголовки
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        _apply_header_style(cell)

    # Записываем данные студентов
    for row_idx, row_data in enumerate(data.grade_rows, 2):
        # №
        cell = ws.cell(row=row_idx, column=1, value=row_idx - 1)
        _apply_cell_style(cell)

        # ФИО
        cell = ws.cell(row=row_idx, column=2, value=row_data.student_name)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = THIN_BORDER

        # Подгруппа
        subgroup_val = row_data.subgroup if row_data.subgroup else "-"
        cell = ws.cell(row=row_idx, column=3, value=subgroup_val)
        _apply_cell_style(cell)

        # Оценки по работам
        col_offset = 4
        for i, key in enumerate(sorted_work_keys):
            grade = row_data.grades_by_work.get(key)
            cell = ws.cell(row=row_idx, column=col_offset + i, value=grade if grade else "")
            _apply_cell_style(cell)

        # Статистика
        stats_col = col_offset + len(sorted_work_keys)

        cell = ws.cell(row=row_idx, column=stats_col, value=row_data.grades_count)
        _apply_cell_style(cell)

        avg_val = f"{row_data.average_grade:.2f}" if row_data.average_grade else "-"
        cell = ws.cell(row=row_idx, column=stats_col + 1, value=avg_val)
        _apply_cell_style(cell)

    _auto_column_width(ws)
    ws.column_dimensions["B"].width = 25

    logger.debug("Лист 'Оценки' создан: %d строк", len(data.grade_rows))


def generate_summary_sheet(wb: Workbook, data: JournalExportData) -> None:
    """
    Создаёт лист "Сводка".

    Args:
        wb: Рабочая книга Excel
        data: Данные журнала для экспорта
    """
    ws = wb.create_sheet("Сводка")

    # Информация о группе
    info_rows = [
        ("Группа:", data.meta.group_code),
        ("Название:", data.meta.group_name or "-"),
        ("Период:", f"{data.meta.period_start.strftime('%d.%m.%Y')} - {data.meta.period_end.strftime('%d.%m.%Y')}"),
        ("Дата генерации:", data.meta.generated_at.strftime("%d.%m.%Y %H:%M")),
        ("", ""),
        ("Статистика:", ""),
        ("Всего студентов:", data.meta.total_students),
        ("Всего занятий:", data.meta.total_lessons),
    ]

    # Средняя посещаемость группы
    if data.attendance_rows:
        avg_attendance = sum(r.attendance_rate for r in data.attendance_rows) / len(data.attendance_rows)
        info_rows.append(("Средняя посещаемость:", f"{avg_attendance:.1f}%"))

    # Средняя оценка группы
    grades_with_avg = [r for r in data.grade_rows if r.average_grade is not None]
    if grades_with_avg:
        avg_grade = sum(r.average_grade for r in grades_with_avg if r.average_grade is not None) / len(grades_with_avg)
        info_rows.append(("Средняя оценка:", f"{avg_grade:.2f}"))

    # Записываем данные
    for row_idx, (label, value) in enumerate(info_rows, 1):
        cell_label = ws.cell(row=row_idx, column=1, value=label)
        cell_label.font = Font(bold=True) if label else Font()
        cell_label.alignment = Alignment(horizontal="left")

        cell_value = ws.cell(row=row_idx, column=2, value=value)
        cell_value.alignment = Alignment(horizontal="left")

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 30

    logger.debug("Лист 'Сводка' создан")


def generate_excel(data: JournalExportData) -> bytes:
    """
    Генерирует Excel файл с данными журнала.

    Args:
        data: Данные журнала для экспорта

    Returns:
        Содержимое Excel файла в байтах
    """
    logger.info(
        "Генерация Excel: группа=%s, студентов=%d, занятий=%d",
        data.meta.group_code,
        data.meta.total_students,
        data.meta.total_lessons,
    )

    wb = Workbook()
    # Удаляем дефолтный лист
    wb.remove(wb.active)

    # Создаём листы
    if data.attendance_rows:
        generate_attendance_sheet(wb, data)

    if data.grade_rows:
        generate_grades_sheet(wb, data)

    generate_summary_sheet(wb, data)

    # Сохраняем в BytesIO
    output = BytesIO()
    wb.save(output)
    content = output.getvalue()

    logger.info("Excel файл сгенерирован: %d байт", len(content))
    return content
